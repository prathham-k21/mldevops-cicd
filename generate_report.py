import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report_docx(filename="LAB-ISE2-CICD_pipeline_and_Monitoring_Tools.docx"):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Header title small
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = header_para.add_run("CICD pipeline and Monitoring Tools\n")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("CICD pipeline and Monitoring Tools\nSummary Report")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph()

    # Table 1: Student Entry Details
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.2), Inches(4.3)]
    headers = [
        ("Details", "Student Entry"),
        ("Student Name", "Pratham Kadam"),
        ("Roll Number(s)", "DS 14 (PRN: 2324000362)"),
        ("Department / Class", "Final Year B.Tech CSE (Data Science) - Semester VII"),
        ("Academic Year", "A.Y 2026-27"),
        ("GitHub Repository URL", "https://github.com/prathham-k21/mldevops-cicd.git")
    ]

    for i, (col1, col2) in enumerate(headers):
        row = table.rows[i]
        cell_0 = row.cells[0]
        cell_1 = row.cells[1]
        
        cell_0.width = col_widths[0]
        cell_1.width = col_widths[1]
        
        cell_0.text = col1
        cell_1.text = col2
        
        if i == 0:
            set_cell_background(cell_0, "17476E")
            set_cell_background(cell_1, "17476E")
            cell_0.paragraphs[0].runs[0].font.bold = True
            cell_0.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            cell_1.paragraphs[0].runs[0].font.bold = True
            cell_1.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        else:
            cell_0.paragraphs[0].runs[0].font.bold = False
            cell_0.paragraphs[0].runs[0].font.size = Pt(10.5)
            if len(cell_1.paragraphs[0].runs) > 0:
                cell_1.paragraphs[0].runs[0].font.size = Pt(10.5)
                cell_1.paragraphs[0].runs[0].font.bold = True
                cell_1.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 45, 98)

    doc.add_page_break()

    # Section 1: Problem Statement
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Problem Statement")
    r1.font.name = "Calibri"
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(31, 78, 121)

    p_ps = doc.add_paragraph()
    p_ps.paragraph_format.line_spacing = 1.15
    p_ps.paragraph_format.space_after = Pt(12)
    run_ps = p_ps.add_run(
        "Develop an end-to-end Machine Learning Web Application for Heart Disease Risk Prediction (CardioPulse) "
        "and implement the complete DevOps / MLOps lifecycle on a Windows PC. The project utilizes Scikit-Learn for "
        "training an accurate clinical risk classification model and Flask for serving both a responsive web interface "
        "and a REST API. Git is employed for local version control, GitHub for cloud source-code hosting, Docker for "
        "application containerization, and Jenkins for automating the continuous integration and deployment pipeline "
        "(including automated pytest suites and image builds). In addition, Prometheus monitoring and health-check "
        "endpoints are incorporated to capture real-time latency, request rates, and prediction distribution metrics. "
        "Each tool is installed, configured, verified, and orchestrated using Windows commands to ensure seamless "
        "transition from development to production."
    )
    run_ps.font.name = "Calibri"
    run_ps.font.size = Pt(11)

    # Section 2: Tools: Environment Setup
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Tools: Environment Setup")
    r2.font.name = "Calibri"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(31, 78, 121)

    # Sub-tool helper
    def add_tool_section(num_str, title, env_bullets, cmds):
        sub_h = doc.add_paragraph()
        sub_h.paragraph_format.space_before = Pt(8)
        sub_h.paragraph_format.space_after = Pt(4)
        run_st = sub_h.add_run(f"{num_str}. {title}")
        run_st.font.name = "Calibri"
        run_st.font.size = Pt(12)
        run_st.font.bold = True

        p_env = doc.add_paragraph()
        r_env = p_env.add_run("1. Environment Setup")
        r_env.font.name = "Calibri"
        r_env.font.bold = True
        r_env.font.size = Pt(11)
        p_env.paragraph_format.space_after = Pt(2)

        for b in env_bullets:
            bp = doc.add_paragraph(style='List Bullet')
            br = bp.add_run(b)
            br.font.name = "Calibri"
            br.font.size = Pt(10.5)
            bp.paragraph_format.space_after = Pt(2)

        p_cmd = doc.add_paragraph()
        r_cmd = p_cmd.add_run("2. Commands to Run on Windows")
        r_cmd.font.name = "Calibri"
        r_cmd.font.bold = True
        r_cmd.font.size = Pt(11)
        p_cmd.paragraph_format.space_before = Pt(4)
        p_cmd.paragraph_format.space_after = Pt(2)

        for c in cmds:
            cp = doc.add_paragraph(style='List Bullet')
            cr = cp.add_run(c)
            cr.font.name = "Consolas"
            cr.font.size = Pt(10)
            cp.paragraph_format.space_after = Pt(2)

    # 1. Python and Flask
    add_tool_section(
        "1", "Python and Flask — Web Application & Machine Learning Development",
        [
            "Install Python 3.13 and configure the system PATH environment variable.",
            "Install Flask, Scikit-Learn, Pandas, Joblib, Prometheus Client, and Pytest.",
            "Train and validate the Heart Disease Random Forest model via train.py, generating model/heart_disease_model.joblib.",
            "Verify installed versions and launch the Flask web application."
        ],
        [
            "python --version",
            "pip --version",
            "pip install -r requirements.txt",
            "python train.py",
            "python -m pytest -v",
            "python app.py"
        ]
    )

    # 2. Git
    add_tool_section(
        "2", "Git — Version Control",
        [
            "Install Git for Windows and verify CLI availability.",
            "Configure the Git username and user email globally.",
            "Initialize local git repository and stage all MLOps application files."
        ],
        [
            "git --version",
            "git init",
            "git status",
            "git add .",
            "git commit -m \"Initial commit: CardioPulse ML App with CI/CD & Monitoring\""
        ]
    )

    # 3. GitHub
    add_tool_section(
        "3", "GitHub — Source-Code Hosting",
        [
            "Create a GitHub account and a remote repository for the project.",
            "Configure Git authentication using Git Credential Manager or Personal Access Token (PAT).",
            "Set the default branch to main and push local commits to the remote origin."
        ],
        [
            "git remote add origin REPOSITORY_URL",
            "git branch -M main",
            "git push -u origin main",
            "git pull origin main"
        ]
    )

    # 4. Docker
    add_tool_section(
        "4", "Docker — Application Containerization",
        [
            "Install Docker Desktop on Windows and configure WSL 2 backend.",
            "Verify Docker Engine daemon and Docker Compose availability.",
            "Build the container image using the Dockerfile and launch the container on port 5000."
        ],
        [
            "docker --version",
            "docker compose version",
            "docker info",
            "docker build -t cardiopulse-ml-app .",
            "docker run -d -p 5000:5000 --name cardiopulse-service cardiopulse-ml-app",
            "docker ps"
        ]
    )

    # 5. Jenkins
    add_tool_section(
        "5", "Jenkins — Continuous Integration and Continuous Deployment",
        [
            "Verify Java JDK 23 compatibility with Jenkins.",
            "Install and launch Jenkins service (or execute Jenkins Docker container).",
            "Configure Jenkins Declarative Pipeline job linked to the GitHub repository and execute Jenkinsfile."
        ],
        [
            "For Jenkins installed directly on Windows:",
            "  java -version",
            "  git --version",
            "  docker --version",
            "  jenkins.exe start (if running as Windows Service)",
            "For Jenkins running as a Docker container named jenkins:",
            "  docker run -d -p 8080:8080 -p 50000:50000 --name jenkins jenkins/jenkins:lts",
            "  docker ps",
            "  docker logs jenkins",
            "  docker exec jenkins java -version"
        ]
    )

    # 6. Monitoring Tools
    add_tool_section(
        "6", "Monitoring & Telemetry — Prometheus & Health Endpoints",
        [
            "Integrate prometheus-client into Flask application to collect real-time HTTP metrics and model inference counters.",
            "Configure health probe endpoint (/health) for container liveness and readiness verification.",
            "Expose Prometheus telemetry endpoint (/metrics) to track latency, request counts, and prediction drift."
        ],
        [
            "curl -s http://localhost:5000/health",
            "curl -s http://localhost:5000/metrics",
            "curl -X POST http://localhost:5000/predict -H \"Content-Type: application/json\" -d \"{\\\"age\\\":55,\\\"sex\\\":1,\\\"cp\\\":0,\\\"trestbps\\\":140,\\\"chol\\\":250,\\\"fbs\\\":0,\\\"restecg\\\":1,\\\"thalach\\\":140,\\\"exang\\\":1,\\\"oldpeak\\\":2.0}\""
        ]
    )

    doc.save(filename)
    print(f"Successfully generated {filename}")

if __name__ == "__main__":
    create_report_docx()

